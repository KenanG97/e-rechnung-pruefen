"""Mitgliederliste aus CSV oder Excel (.xlsx) lesen – Spalten werden an ihren Namen erkannt."""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from typing import Optional

from kugrenze.bankcsv import parse_amount, parse_day

from .core import Debit

COLS = {
    "name": ["name", "kontoinhaber", "zahler", "mitglied", "vor- und nachname", "vollständiger name", "debitor"],
    "first": ["vorname", "first name"],
    "last": ["nachname", "familienname", "last name", "zuname"],
    "iban": ["iban"],
    "bic": ["bic", "swift"],
    "amount": ["betrag", "beitrag", "mitgliedsbeitrag", "summe", "amount", "jahresbeitrag"],
    "mandate_id": ["mandatsreferenz", "mandats-referenz", "mandat", "mandatsnummer", "mandate", "mandat-id", "mandatsid"],
    "mandate_date": ["mandatsdatum", "datum der unterschrift", "unterschriftsdatum", "mandat vom", "datum mandat", "signature date",
                     "mandat datum", "datum"],
    "purpose": ["verwendungszweck", "zweck", "text", "purpose"],
    "street": ["straße", "strasse", "adresse", "anschrift", "street"],
    "postcode": ["plz", "postleitzahl", "zip"],
    "town": ["ort", "stadt", "wohnort", "city", "town"],
    "country": ["land", "country", "ländercode"],
    "email": ["e-mail", "email", "mail"],
    "seq": ["sequenz", "lastschriftart", "sequence"],
}


class MemberError(Exception):
    pass


def _rows(data: bytes, filename: str) -> list[list[str]]:
    if filename.lower().endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        try:
            from openpyxl import load_workbook
        except ImportError as e:  # pragma: no cover
            raise MemberError("Excel-Unterstützung nicht verfügbar – bitte als CSV speichern.") from e
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:
            raise MemberError(f"Excel-Datei nicht lesbar: {e}") from e
        ws = wb.worksheets[0]
        out = []
        for r in ws.iter_rows(values_only=True):
            out.append(["" if v is None else (v.strftime("%d.%m.%Y") if isinstance(v, (datetime, date)) else
                                              (f"{v:.2f}" if isinstance(v, float) else str(v))) for v in r])
        return out
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise MemberError("Zeichenkodierung nicht erkannt.")
    first = text.split("\n", 1)[0]
    delim = max((";", ",", "\t"), key=first.count)
    return list(csv.reader(io.StringIO(text), delimiter=delim))


def _find(hn: list[str], keys: list[str], used: set[int]) -> Optional[int]:
    for k in keys:
        for i, h in enumerate(hn):
            if i not in used and h == k:
                return i
    for k in keys:
        for i, h in enumerate(hn):
            if i not in used and k in h:
                return i
    return None


def load_members(data: bytes, filename: str = "", default_amount: str = "", default_purpose: str = "") -> tuple[list[Debit], dict]:
    rows = _rows(data, filename)
    hidx = None
    for i, r in enumerate(rows[:15]):
        if any("iban" in (c or "").lower() for c in r):
            hidx = i
            break
    if hidx is None:
        raise MemberError("Keine Spalte „IBAN“ gefunden. Die erste Zeile der Liste braucht Spaltenüberschriften.")
    hn = [(c or "").strip().lower() for c in rows[hidx]]
    used: set[int] = set()
    cols = {}
    for k in ("iban", "bic", "mandate_id", "mandate_date", "amount", "first", "last", "name", "purpose", "street",
              "postcode", "town", "country", "email", "seq"):
        c = _find(hn, COLS[k], used)
        cols[k] = c
        if c is not None:
            used.add(c)
    if cols["name"] is None and cols["last"] is None:
        raise MemberError("Keine Spalte mit Namen gefunden (z. B. „Name“ oder „Vorname“/„Nachname“).")
    out = []
    for n, r in enumerate(rows[hidx + 1:], start=hidx + 2):
        g = lambda k: (str(r[cols[k]]).strip() if cols[k] is not None and cols[k] < len(r) and r[cols[k]] is not None else "")  # noqa: E731
        if not any(str(c).strip() for c in r):
            continue
        name = g("name") or " ".join(x for x in (g("first"), g("last")) if x)
        amt_raw = g("amount") or default_amount
        seq = g("seq").upper()
        seq = {"ERST": "FRST", "ERSTLASTSCHRIFT": "FRST", "EINMALIG": "OOFF", "LETZTE": "FNAL", "FOLGE": "RCUR"}.get(seq, seq)
        out.append(Debit(
            row=n, name=name, iban=g("iban"), bic=g("bic"), amount=parse_amount(amt_raw) if amt_raw else None,
            mandate_id=g("mandate_id"), mandate_date=parse_day(g("mandate_date")),
            purpose=g("purpose") or default_purpose, street=g("street"), postcode=g("postcode"), town=g("town"),
            country=(g("country") if len(g("country")) == 2 else {"österreich": "AT", "oesterreich": "AT", "deutschland": "DE",
                     "schweiz": "CH", "italien": "IT", "liechtenstein": "LI"}.get(g("country").lower(), g("country"))).upper(),
            email=g("email"), seq=seq if seq in ("FRST", "RCUR", "OOFF", "FNAL") else "RCUR",
        ))
    if not out:
        raise MemberError("Die Liste enthält keine Mitglieder.")
    mapping = {k: rows[hidx][v] for k, v in cols.items() if v is not None}
    return out, mapping
